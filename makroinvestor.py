#!/usr/bin/env python3
"""
Makroinvestor — analyserer makrotendenser og anbefaler sektorer at investere i.
Baseret på Excel-arket: Makroinvestoren_Final.xlsx
"""

import sys
import os
import argparse
import openpyxl

# ── ANSI farver ──────────────────────────────────────────────────────────────
RESET  = "\033[0m"
BOLD   = "\033[1m"
GREEN  = "\033[92m"
YELLOW = "\033[93m"
RED    = "\033[91m"
CYAN   = "\033[96m"
BLUE   = "\033[94m"
GRAY   = "\033[90m"
WHITE  = "\033[97m"

DEFAULT_EXCEL = os.path.join(os.path.dirname(__file__),
    "/root/.claude/uploads/4353b6a9-163f-529f-a054-f6c3b69a4fdb/4eb3c991-Makroinvestoren_Final.xlsx")

# ── Hjælpefunktioner ─────────────────────────────────────────────────────────

def c(color, text):
    return f"{color}{text}{RESET}"

def header(title):
    w = 62
    print()
    print(c(CYAN, "═" * w))
    print(c(CYAN, f"  {BOLD}{title}"))
    print(c(CYAN, "═" * w))

def subheader(title):
    print(f"\n{c(BOLD, '── ' + title + ' ──')}")

def bar(score, max_score=5.0, width=20):
    filled = int(round(score / max_score * width))
    b = "█" * filled + "░" * (width - filled)
    if score >= 3.5:
        col = GREEN
    elif score >= 2.5:
        col = YELLOW
    else:
        col = RED
    return c(col, b) + f" {score:.2f}"

def phase_color(phase):
    colors = {"Early": GREEN, "Mid": BLUE, "Late": YELLOW, "Recession": RED}
    return c(colors.get(phase, WHITE), phase)

def ask(question, options):
    """Stil et spørgsmål med nummererede svar og returnér valgt score."""
    print(f"\n{c(BOLD, question)}")
    for i, (opt, _score) in enumerate(options, 1):
        print(f"  {c(GRAY, str(i)+'.')} {opt}")
    while True:
        try:
            val = int(input(c(CYAN, "  Dit svar (1-4): ")).strip())
            if 1 <= val <= len(options):
                return options[val - 1][1]
        except (ValueError, EOFError):
            pass
        print(c(RED, "  Indtast venligst et tal mellem 1 og 4."))

# ── Indlæs Excel ─────────────────────────────────────────────────────────────

def load_workbook(path):
    try:
        return openpyxl.load_workbook(path, data_only=True)
    except FileNotFoundError:
        print(c(RED, f"\nFejl: Kan ikke finde Excel-filen: {path}"))
        print("Brug --excel <sti> til at angive en anden sti.")
        sys.exit(1)

def parse_rangliste(ws):
    """Returnér dict sektor -> (Q1_2026, Q2_2026)"""
    result = {}
    for row in ws.iter_rows(values_only=True):
        sektor = row[2] if len(row) > 2 else None
        q1 = row[3] if len(row) > 3 else None
        q2 = row[4] if len(row) > 4 else None
        if isinstance(sektor, str) and sektor not in ("Sektor",) and isinstance(q1, (int, float)):
            result[sektor] = (float(q1), float(q2) if isinstance(q2, (int, float)) else float(q1))
    return result

def parse_bnp_phases(ws):
    """Returnér dict sektor -> {fase -> int} og BNP-faser per region."""
    phase_scores = {"++": 2, "+": 1, None: 0, "–": -1, "--": -2, "-": -1}
    sectors = {}
    bnp_phases = {"Danmark": [], "Europa": [], "USA": []}  # (kvartal, fase)

    rows = list(ws.iter_rows(values_only=True))

    # Sektor-fase sensitivitet (rækker 2-13, 0-indeks)
    phase_order = ["Early", "Mid", "Late", "Recession"]
    for row in rows[2:14]:
        sektor = row[1]
        if isinstance(sektor, str):
            entry = {}
            for i, phase in enumerate(phase_order):
                val = row[2 + i] if len(row) > 2 + i else None
                entry[phase] = phase_scores.get(val, 0)
            sectors[sektor.strip()] = entry

    # BNP-historik per region
    current_region = None
    for row in rows:
        if row[1] in ("Danmark", "Europa", "USA") and row[2] is None:
            current_region = row[1]
        if current_region and isinstance(row[2], str) and row[2].startswith("Q") and isinstance(row[4], str):
            bnp_phases[current_region].append((row[2], row[4]))

    return sectors, bnp_phases

def parse_pmi(ws):
    """Returnér seneste PMI-fase per region."""
    regions = {"USA PMI": "USA", "Euroområdet PMI": "Europa", "Danmark PMI": "Danmark"}
    latest = {}
    current_header = None
    for row in ws.iter_rows(values_only=True):
        label = row[4] if len(row) > 4 else None
        if label in regions:
            current_header = regions[label]
        if current_header and isinstance(row[3], str) and row[3].startswith("Q") and isinstance(row[5], str):
            latest[current_header] = (row[3], row[5])
    return latest

def parse_macro_indicators(wb):
    """Hent seneste fase-signaler fra alle makroindikatorer."""
    signals = {}

    def last_phase(sheetname, region_col, phase_col, region_label, header_col=4):
        if sheetname not in wb.sheetnames:
            return None
        ws = wb[sheetname]
        current = None
        last = None
        for row in ws.iter_rows(values_only=True):
            lbl = row[header_col] if len(row) > header_col else None
            if lbl == region_label:
                current = True
            if current and isinstance(row[region_col], str) and row[region_col].startswith("Q") and isinstance(row[phase_col], str):
                last = (row[region_col], row[phase_col])
        return last

    # PMI
    pmi = parse_pmi(wb["PMI"])
    signals["PMI"] = pmi

    # 10yr rate
    rate_phases = {}
    ws = wb["10 yr rate"]
    mapping = {3: ("USA 10 year", "USA"), 9: ("Euroområdet 10 year", "Europa")}
    for col_kv, (label, region) in mapping.items():
        last = None
        for row in ws.iter_rows(values_only=True):
            if row[col_kv] is not None and isinstance(row[col_kv], str) and row[col_kv].startswith("Q"):
                if len(row) > col_kv + 2 and isinstance(row[col_kv + 2], str):
                    last = (row[col_kv], row[col_kv + 2])
        if last:
            rate_phases[region] = last
    # Simple parse: scan all rows
    rate_phases2 = {}
    ws = wb["10 yr rate"]
    cur_region = None
    for row in ws.iter_rows(values_only=True):
        if row[2] == "Kvartal" and row[3] == "USA 10 year":
            cur_region = "USA"
        elif row[2] == "Kvartal" and row[3] == "Euroområdet 10 year":
            cur_region = "Europa"
        if cur_region and isinstance(row[2], str) and row[2].startswith("Q") and isinstance(row[4], str):
            rate_phases2[cur_region] = (row[2], row[4])
    signals["10yr"] = rate_phases2

    # CPI
    cpi_phases = {}
    ws = wb["CPI(Inflation)"]
    cur_region = None
    for row in ws.iter_rows(values_only=True):
        if row[2] == "Kvartal":
            cur_region = str(row[3]).split()[0] if row[3] else None
            if "Danmark" in str(row[3]):
                cur_region = "Danmark"
            elif "Euroområdet" in str(row[3]):
                cur_region = "Europa"
        if cur_region and isinstance(row[2], str) and row[2].startswith("Q") and isinstance(row[4], str):
            cpi_phases[cur_region] = (row[2], row[4])
    signals["CPI"] = cpi_phases

    # VIX
    vix_phases = {}
    ws = wb["VIX"]
    cur_region = None
    for row in ws.iter_rows(values_only=True):
        if row[3] == "Kvartal" and row[4] == "Europa":
            cur_region = "Europa"
        elif row[3] == "Kvartal" and row[4] == "USA":
            cur_region = "USA"
        if cur_region and isinstance(row[3], str) and row[3].startswith("Q") and isinstance(row[5], str):
            vix_phases[cur_region] = (row[3], row[5])
    signals["VIX"] = vix_phases

    return signals

def determine_phase(signals, region="Europa"):
    """Bestem den dominerende fase for en region ud fra alle signaler."""
    phase_points = {"Early": 0, "Mid": 0, "Late": 0, "Recession": 0}
    sources = []

    for indicator, data in signals.items():
        if indicator == "PMI":
            entry = data.get(region)
        else:
            entry = data.get(region)

        if entry:
            kvartal, fase = entry
            if fase in phase_points:
                phase_points[fase] += 1
                sources.append((indicator, kvartal, fase))

    if not sources:
        return "Mid", []

    dominant = max(phase_points, key=phase_points.get)
    return dominant, sources

def parse_etf_liste(ws):
    """Returnér dict: sektor (normalized) -> liste af ETF dicts."""
    etfs = {}
    for row in ws.iter_rows(values_only=True):
        if row[0] == "Kategori":
            continue
        kategori, navn, isin, region, sektor = (row[i] if len(row) > i else None for i in range(5))
        if navn and sektor:
            key = sektor.strip().lower()
            if key not in etfs:
                etfs[key] = []
            etfs[key].append({"navn": navn, "isin": isin, "region": region, "sektor": sektor})
    return etfs

def parse_aktieliste(ws):
    """Returnér liste af aktie-dicts, filtreret for gyldige rækker."""
    aktier = []
    header_done = False
    for row in ws.iter_rows(values_only=True):
        if not header_done:
            header_done = True
            continue
        navn = row[0] if row[0] else None
        price = row[1]
        mkt_cap = row[2]
        sektor = row[5] if len(row) > 5 else None
        ticker = row[7] if len(row) > 7 else None
        pe = row[6] if len(row) > 6 else None
        beta = row[9] if len(row) > 9 else None

        if ticker and isinstance(ticker, str) and sektor and isinstance(mkt_cap, (int, float)):
            aktier.append({
                "ticker": ticker,
                "sektor": sektor.strip() if sektor else "",
                "price": price,
                "market_cap": mkt_cap,
                "pe": pe,
                "beta": beta,
            })
    return aktier

def parse_sporgeskema(ws):
    """Returnér liste af (question, [(svar, score), ...])."""
    questions = []
    for row in ws.iter_rows(values_only=True):
        q = row[0]
        if not isinstance(q, str) or q.startswith("Spørg"):
            continue
        svar = [row[i] for i in range(1, 5) if row[i] is not None]
        scores = [row[i] for i in range(5, 9) if row[i] is not None]
        pairs = list(zip(svar, scores))
        if pairs:
            questions.append((q, pairs))
    return questions

def parse_profiler(ws):
    """Returnér liste af profil-dicts."""
    profiler = []
    for row in ws.iter_rows(values_only=True):
        if row[0] == "Profilnavn":
            continue
        navn, fra, til, beskrivelse, aktivfordeling, alternativer = (row[i] if len(row) > i else None for i in range(6))
        if navn and isinstance(fra, (int, float)):
            profiler.append({
                "navn": navn, "fra": fra, "til": til,
                "beskrivelse": beskrivelse,
                "aktivfordeling": aktivfordeling,
                "alternativer": alternativer,
            })
    return profiler

# ── Sektornormalisering ───────────────────────────────────────────────────────

SEKTOR_ALIAS = {
    "infomation technology": "Information Technology",
    "information technology": "Information Technology",
    "financials": "Financials",
    "real estate": "Real Estate",
    "consumer discretionary": "Consumer Discretionary",
    "industrials": "Industrials",
    "materials": "Materials",
    "consumer staples": "Consumer Staples",
    "health care": "Health Care",
    "healthcare": "Health Care",
    "energy": "Energy",
    "communications sercives": "Communications Services",
    "communications services": "Communications Services",
    "communication services": "Communications Services",
    "utilities": "Utilities",
}

def norm_sektor(s):
    if not s:
        return s
    return SEKTOR_ALIAS.get(s.strip().lower(), s.strip())

# ── Beregn sektor-scores ──────────────────────────────────────────────────────

def compute_sector_scores(dk_scores, eu_scores, usa_scores, bnp_sectors, phase):
    """Kombiner makroscorer og BNP-fase til en samlet score per sektor."""
    all_sektorer = set(dk_scores) | set(eu_scores) | set(usa_scores)
    results = {}

    for sektor in all_sektorer:
        region_scores = []
        for d in [dk_scores, eu_scores, usa_scores]:
            if sektor in d:
                region_scores.append(d[sektor][1])  # Q2 2026

        if not region_scores:
            continue

        makro_avg = sum(region_scores) / len(region_scores)

        # Find BNP-bonus (match på normaliseret sektornavn)
        bnp_bonus = 0
        for bnp_key, phases in bnp_sectors.items():
            if norm_sektor(bnp_key) == sektor:
                bnp_bonus = phases.get(phase, 0)
                break

        # Kombiner: 70% makroscore (skala 1-5) + 30% BNP-bonus (skala -2 til +2 -> 0-5)
        bnp_normalized = (bnp_bonus + 2) / 4 * 5  # omregn til 0-5 skala
        composite = makro_avg * 0.70 + bnp_normalized * 0.30

        results[sektor] = {
            "makro_avg": makro_avg,
            "bnp_bonus": bnp_bonus,
            "composite": composite,
            "dk": dk_scores.get(sektor, (None, None))[1],
            "eu": eu_scores.get(sektor, (None, None))[1],
            "usa": usa_scores.get(sektor, (None, None))[1],
        }

    return dict(sorted(results.items(), key=lambda x: x[1]["composite"], reverse=True))

# ── Risikoprofil ──────────────────────────────────────────────────────────────

def run_questionnaire(sporgeskema, profiler):
    header("Risikoprofil — 8 spørgsmål")
    print(c(GRAY, "  Besvar hvert spørgsmål ved at taste et tal (1-4)."))

    total_score = 0
    max_score = 0

    for q, pairs in sporgeskema:
        score = ask(q, pairs)
        total_score += score
        max_score += max(p[1] for p in pairs)

    ratio = total_score / max_score

    # Match profil
    profil = profiler[-1]
    for p in profiler:
        if p["fra"] <= ratio <= p["til"]:
            profil = p
            break

    return profil, ratio, total_score, max_score

# ── Vis resultater ────────────────────────────────────────────────────────────

def vis_makro_analyse(signals, fase_dk, fase_eu, fase_usa):
    header("Makrofaseanalyse")

    for region, fase, sources in [
        ("Danmark", fase_dk[0], fase_dk[1]),
        ("Europa",  fase_eu[0], fase_eu[1]),
        ("USA",     fase_usa[0], fase_usa[1]),
    ]:
        print(f"\n  {c(BOLD, region):20s} → Fase: {phase_color(fase)}")
        for ind, kv, f in sources:
            print(f"    {c(GRAY, f'  {ind} ({kv}):')} {phase_color(f)}")

def vis_sektor_ranking(scored, top_n=5):
    header(f"Sektorer — Top {top_n} anbefalinger")
    print(c(GRAY, "  Baseret på makroscorer (Q2 2026) + BNP-fasebonus\n"))
    print(f"  {'Rang':<5} {'Sektor':<28} {'Score':>6}  {'Bar'}")
    print(f"  {'-'*65}")

    for i, (sektor, data) in enumerate(list(scored.items())[:top_n], 1):
        rank_col = GREEN if i <= 3 else YELLOW
        bnp_str = f"{'+' if data['bnp_bonus'] >= 0 else ''}{data['bnp_bonus']}"
        print(f"  {c(rank_col, f'#{i}')}    {sektor:<28} {data['composite']:5.2f}  {bar(data['composite'])}"
              f"  {c(GRAY, f'BNP:{bnp_str}')}")

    print(f"\n  {c(GRAY, 'Regionsscorer (Q2 2026):')}")
    print(f"  {'Sektor':<28} {'DK':>6} {'EU':>6} {'USA':>6}")
    print(f"  {'-'*50}")
    for sektor, data in list(scored.items())[:top_n]:
        dk  = f"{data['dk']:.2f}"  if data['dk']  else "  N/A"
        eu  = f"{data['eu']:.2f}"  if data['eu']  else "  N/A"
        usa = f"{data['usa']:.2f}" if data['usa'] else "  N/A"
        print(f"  {sektor:<28} {dk:>6} {eu:>6} {usa:>6}")

def vis_etf_og_aktier(scored, etf_liste, aktier_dk, aktier_eu, aktier_usa, top_n=3):
    header("ETF'er & Aktier pr. top-sektor")

    for sektor, data in list(scored.items())[:top_n]:
        subheader(f"{sektor}  (score: {data['composite']:.2f})")

        # ETF'er
        sektor_lower = sektor.lower()
        etf_match = []
        for key, etfs in etf_liste.items():
            if key in sektor_lower or sektor_lower in key:
                etf_match.extend(etfs)
        # fallback — fuzzy
        if not etf_match:
            for key, etfs in etf_liste.items():
                words = sektor_lower.split()
                if any(w in key for w in words if len(w) > 4):
                    etf_match.extend(etfs)

        if etf_match:
            etf_label = c(BOLD, "ETF'er:")
            print(f"  {etf_label}")
            for e in etf_match[:3]:
                print(f"    • {e['navn']}  {c(GRAY, '(' + (e['isin'] or '') + ')')}")
        else:
            print(f"  {c(GRAY, 'Ingen ETF fundet for denne sektor.')}")

        # Aktier
        for region_label, aktier in [("Danmark", aktier_dk), ("Europa", aktier_eu), ("USA", aktier_usa)]:
            matches = [a for a in aktier if norm_sektor(a["sektor"]) == sektor]
            matches.sort(key=lambda a: a["market_cap"] or 0, reverse=True)
            if matches:
                print(f"  {c(BOLD, region_label + ':')} ", end="")
                print(", ".join(f"{a['ticker']}" for a in matches[:3]))

def vis_profil_og_fordeling(profil, beloeb):
    header("Din Risikoprofil & Porteføljefordeling")
    print(f"\n  Profil:       {c(BOLD + GREEN, profil['navn'])}")
    print(f"  Beskrivelse:  {profil['beskrivelse']}")
    print(f"  Alternativer: {c(GRAY, profil['alternativer'] or 'N/A')}")
    print(f"\n  {c(BOLD, 'Anbefalet aktivfordeling:')}")

    fordeling_str = profil["aktivfordeling"] or ""
    print(f"  {fordeling_str}")

    if beloeb > 0:
        print(f"\n  {c(BOLD, f'Fordeling af {beloeb:,.0f} DKK:')}")
        # Parse "Obligationer 50%, ETF 25%, Aktier 15%, Alternative 10%"
        import re
        parts = re.findall(r'([A-Za-zæøåÆØÅ ]+?)\s+(\d+)%', fordeling_str)
        for navn, pct in parts:
            amount = beloeb * int(pct) / 100
            bar_w = int(int(pct) / 5)
            col = GREEN if int(pct) >= 40 else (YELLOW if int(pct) >= 20 else BLUE)
            print(f"    {navn.strip():<20} {pct:>3}%  {c(col, '█' * bar_w):<30} {amount:>12,.0f} DKK")

# ── Hovedprogram ──────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Makroinvestor — sektoranalyse baseret på makrotendenser")
    parser.add_argument("--excel", default=DEFAULT_EXCEL, help="Sti til Excel-filen")
    parser.add_argument("--skip-questionnaire", action="store_true", help="Spring spørgeskema over (bruger Balanceret)")
    args = parser.parse_args()

    print(c(BOLD + CYAN, "\n╔══════════════════════════════════════════════════════════════╗"))
    print(c(BOLD + CYAN,   "║          M A K R O I N V E S T O R  2026                    ║"))
    print(c(BOLD + CYAN,   "╚══════════════════════════════════════════════════════════════╝"))
    print(c(GRAY, "  Analyserer makrotendenser og finder de bedste sektorer.\n"))

    # Indlæs data
    wb = load_workbook(args.excel)

    dk_scores  = parse_rangliste(wb["Ranglisten DK"])
    eu_scores  = parse_rangliste(wb["Ranglisten EU"])
    usa_scores = parse_rangliste(wb["Ranglisten USA"])

    bnp_sectors, _ = parse_bnp_phases(wb["BNP"])
    etf_liste       = parse_etf_liste(wb["ETF - Liste"])
    aktier_dk       = parse_aktieliste(wb["Aktieliste DK"])
    aktier_eu       = parse_aktieliste(wb["Aktieliste EU"])
    aktier_usa      = parse_aktieliste(wb["Aktieliste USA"])
    sporgeskema     = parse_sporgeskema(wb["Spørgeskema"])
    profiler        = parse_profiler(wb["Profiler"])

    # Makrofaseanalyse
    signals = parse_macro_indicators(wb)
    fase_dk  = determine_phase(signals, "Danmark")
    fase_eu  = determine_phase(signals, "Europa")
    fase_usa = determine_phase(signals, "USA")

    # Dominerende global fase (simpelt flertal)
    from collections import Counter
    alle_faser = [fase_dk[0], fase_eu[0], fase_usa[0]]
    global_fase = Counter(alle_faser).most_common(1)[0][0]

    vis_makro_analyse(signals, fase_dk, fase_eu, fase_usa)
    print(f"\n  {c(BOLD, 'Dominerende global fase:')} {phase_color(global_fase)}")

    # Sektor-scoring
    scored = compute_sector_scores(dk_scores, eu_scores, usa_scores, bnp_sectors, global_fase)

    vis_sektor_ranking(scored, top_n=5)
    vis_etf_og_aktier(scored, etf_liste, aktier_dk, aktier_eu, aktier_usa, top_n=3)

    # Risikoprofil
    if args.skip_questionnaire:
        profil = next((p for p in profiler if p["navn"] == "Balanceret"), profiler[1])
        print(c(GRAY, "\n  (Spørgeskema sprunget over — bruger Balanceret)"))
    else:
        profil, ratio, total, max_s = run_questionnaire(sporgeskema, profiler)
        print(f"\n  {c(GRAY, f'Score: {total}/{max_s} ({ratio:.0%})')}")

    # Investeringsbeløb
    header("Investeringsbeløb")
    while True:
        try:
            beloeb_str = input(c(CYAN, "  Hvor meget ønsker du at investere (DKK)? ")).strip().replace(".", "").replace(",", "")
            beloeb = float(beloeb_str)
            break
        except (ValueError, EOFError):
            print(c(RED, "  Indtast venligst et gyldigt beløb."))

    vis_profil_og_fordeling(profil, beloeb)

    # Afslutning
    header("Opsummering")
    print(f"  Global makrofase:   {phase_color(global_fase)}")
    print(f"  Risikoprofil:       {c(BOLD, profil['navn'])}")
    print(f"  Bedste sektor:      {c(GREEN + BOLD, list(scored.keys())[0])}")
    print(f"  #2 sektor:          {c(GREEN, list(scored.keys())[1])}")
    print(f"  #3 sektor:          {c(YELLOW, list(scored.keys())[2])}")
    print()
    print(c(GRAY, "  ⚠  Dette er ikke finansiel rådgivning. Invester altid med omhu."))
    print()

if __name__ == "__main__":
    main()
